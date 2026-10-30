"""Structured, evidence-gated DOCX resume renderer.

A data directory holds three things:

  PROFILE.md          canonical facts about the person (the only source of truth)
  resume_facts.json   registry of fact IDs, each pointing back into PROFILE.md
  specs/*.json        one spec per target role, citing fact IDs for every claim

The renderer owns layout only. ``resume.py build`` owns page-fit tuning and
validation.
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x55, 0x55, 0x55)
FAINT = RGBColor(0x77, 0x77, 0x77)
LINK = RGBColor(0x1F, 0x4E, 0x8C)
RULE = "9A9A9A"

# Contact labels accepted in the PROFILE.md "## Contact" block, mapped to keys.
CONTACT_LABELS = {
    "name": "name",
    "location": "location",
    "phone": "phone",
    "email": "email",
    "email (career)": "email",
    "linkedin": "linkedin",
    "github": "github",
    "website": "website",
    "portfolio": "website",
}
CONTACT_ORDER = ["location", "phone", "email", "linkedin", "github", "website"]

# Built-in section ids, in default layout order, with their default headings.
DEFAULT_SECTION_ORDER = ["summary", "alignment", "experience", "leadership", "projects", "skills", "education"]
SECTION_TITLES = dict(
    zip(
        DEFAULT_SECTION_ORDER,
        ["PROFILE", "ROLE ALIGNMENT", "EXPERIENCE", "LEADERSHIP", "PROJECTS", "SKILLS & TOOLS", "EDUCATION"],
    )
)
# Built-ins skipped when the spec leaves them empty or out.
OPTIONAL_SECTIONS = {"alignment", "leadership", "projects", "skills"}
CUSTOM_ID = re.compile(r"[a-z0-9_-]+")


@dataclass(frozen=True)
class Layout:
    """Default section order and headings. A spec picks one with its ``layout`` key."""

    order: tuple[str, ...]
    titles: dict[str, str]
    technical: bool = False
    line_spacing: float = 1.0
    custom_before: str | None = None  # default slot for custom sections; None means last


LAYOUTS = {
    "classic": Layout(order=tuple(DEFAULT_SECTION_ORDER), titles=SECTION_TITLES, custom_before="education"),
    "technical": Layout(
        order=("education", "skills", "projects", "experience", "leadership"),
        titles={
            "education": "EDUCATION",
            "skills": "TECHNICAL SKILLS",
            "projects": "SELECTED TECHNICAL PROJECTS",
            "experience": "PROFESSIONAL EXPERIENCE",
            "leadership": "LEADERSHIP & ACTIVITIES",
        },
        technical=True,
        line_spacing=1.06,
    ),
}


def layout_of(spec: dict[str, Any]) -> Layout:
    name = spec.get("layout", "classic")
    if name not in LAYOUTS:
        raise ValueError(f"unknown layout {name!r}; expected one of {sorted(LAYOUTS)}")
    return LAYOUTS[name]


@dataclass(frozen=True)
class Style:
    """Readable typography plus an independently tunable spacing scale."""

    body_pt: float = 10.5
    spacing_scale: float = 1.0
    margin_in: float = 0.5
    line_spacing: float = 1.0
    font_name: str = "Arial"
    accent_hex: str = "1A56DB"
    technical: bool = False

    @property
    def accent(self) -> RGBColor:
        return RGBColor.from_string(self.accent_hex.upper())

    def font(self, offset: float = 0.0) -> Pt:
        return Pt(self.body_pt + offset)

    def gap(self, points: float) -> Pt:
        return Pt(points * self.spacing_scale)


@dataclass(frozen=True)
class ResumeInputs:
    spec: dict[str, Any]
    facts: dict[str, Any]
    person: dict[str, str]
    spec_path: Path
    facts_path: Path
    profile_path: Path

    def base_style(self) -> Style:
        theme = {**self.facts.get("style", {}), **self.spec.get("style", {})}
        return Style(
            font_name=theme.get("font", Style.font_name),
            accent_hex=theme.get("accent", Style.accent_hex).lstrip("#"),
        )


def load_inputs(spec_path: Path, facts_path: Path) -> ResumeInputs:
    spec_path = Path(spec_path).resolve()
    facts_path = Path(facts_path).resolve()
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    facts = json.loads(facts_path.read_text(encoding="utf-8"))
    profile_path = _profile_path(facts, facts_path)
    person = _resolve_person(facts, profile_path)
    validate_spec(spec, facts, spec_path)
    return ResumeInputs(
        spec=spec,
        facts=facts,
        person=person,
        spec_path=spec_path,
        facts_path=facts_path,
        profile_path=profile_path,
    )


def _profile_path(facts: dict[str, Any], facts_path: Path) -> Path:
    canonical = Path(facts.get("canonical_profile", "PROFILE.md"))
    if not canonical.is_absolute():
        canonical = (facts_path.parent / canonical).resolve()
    if not canonical.exists():
        raise FileNotFoundError(f"canonical profile not found: {canonical}")
    return canonical


def _normalize_url(value: str) -> str:
    return value.strip().removeprefix("https://").removeprefix("http://").removeprefix("www.").rstrip("/")


def profile_contact(profile_path: Path) -> dict[str, str]:
    """Read identity fields from the ``## Contact`` block of PROFILE.md."""
    parsed: dict[str, str] = {}
    in_contact = False
    for line in profile_path.read_text(encoding="utf-8").splitlines():
        if line.startswith("## "):
            in_contact = line[3:].strip().lower() == "contact"
            continue
        if not in_contact or not line.startswith("- ") or ":" not in line:
            continue
        label, value = line[2:].split(":", 1)
        key = CONTACT_LABELS.get(label.strip().lower())
        if key and value.strip() and key not in parsed:
            parsed[key] = _normalize_url(value)
    return parsed


def _resolve_person(facts: dict[str, Any], profile_path: Path) -> dict[str, str]:
    """PROFILE.md is authoritative; an optional facts ``person`` block must agree."""
    person = profile_contact(profile_path)
    for key in ("name", "email"):
        if not person.get(key):
            raise ValueError(
                f"PROFILE.md ## Contact block is missing '- {key.title()}: ...' ({profile_path})"
            )
    declared = {k: _normalize_url(str(v)) for k, v in facts.get("person", {}).items()}
    drift = {k: (person.get(k), v) for k, v in declared.items() if person.get(k) != v}
    if drift:
        raise ValueError(
            "resume_facts.json person block disagrees with PROFILE.md "
            f"(profile, facts): {drift}. Fix one, or delete the person block."
        )
    return person


def contact_fields(spec: dict[str, Any], person: dict[str, str]) -> dict[str, str]:
    """Contact values that will render for this spec, in display order."""
    hidden = set(spec.get("hide_contact", []))
    if spec.get("hide_phone"):
        hidden.add("phone")
    return {k: person[k] for k in CONTACT_ORDER if person.get(k) and k not in hidden}


def contact_links(person: dict[str, str]) -> dict[str, str]:
    """Link target for each contact field that has one."""
    targets = {
        "email": "mailto:{}",
        "linkedin": "https://www.{}/",
        "github": "https://{}",
        "website": "https://{}",
    }
    return {k: target.format(person[k]) for k, target in targets.items() if person.get(k)}


def expected_links(inputs: ResumeInputs) -> list[str]:
    """Every URL the rendered resume must carry as a clickable link."""
    spec = inputs.spec
    if not layout_of(spec).technical:
        return []
    targets = contact_links(inputs.person)
    links = [targets[k] for k in contact_fields(spec, inputs.person) if k in targets]
    links.extend(link["url"] for item in _linkable_items(spec) for link in item.get("links", []))
    return links


def _linkable_items(spec: dict[str, Any]):
    """Items whose ``links`` render in the technical layout."""
    for section in ("experience", "leadership", "projects"):
        yield from spec.get(section, [])
    for section in spec.get("sections", []):
        if section.get("type") == "entries":
            yield from section.get("items", [])


def _education_entries(spec: dict[str, Any]) -> list[dict[str, Any]]:
    edu = spec.get("education", [])
    return [edu] if isinstance(edu, dict) else list(edu)


def _all_evidence_ids(spec: dict[str, Any]) -> set[str]:
    ids: set[str] = set(spec.get("summary_evidence", []))
    ids.update(spec.get("projects_note_evidence", []))
    for section in ("experience", "leadership", "projects"):
        for item in spec.get(section, []):
            ids.update(item.get("evidence", []))
            for bullet in item.get("bullets", []):
                ids.update(bullet.get("evidence", []))
    for group in spec.get("skills", []):
        ids.update(group.get("evidence", []))
    for edu in _education_entries(spec):
        ids.update(edu.get("evidence", []))
    for section in spec.get("sections", []):
        for item in section.get("items", []):
            ids.update(item.get("evidence", []))
            for bullet in item.get("bullets", []):
                ids.update(bullet.get("evidence", []))
    return ids


def _iter_text(value: Any):
    if isinstance(value, str):
        yield value
    elif isinstance(value, dict):
        for nested in value.values():
            yield from _iter_text(nested)
    elif isinstance(value, list):
        for nested in value:
            yield from _iter_text(nested)


def validate_spec(spec: dict[str, Any], facts: dict[str, Any], spec_path: Path) -> None:
    layout = _resolve_sections(spec)
    has_summary = any(section["id"] == "summary" for section in layout)
    required = {"company", "role", "experience", "education"} | ({"summary"} if has_summary else set())
    missing = sorted(required - set(spec))
    if missing:
        raise ValueError(f"spec missing required keys: {missing}")

    known = set(facts.get("facts", {}))
    unknown = sorted(_all_evidence_ids(spec) - known)
    if unknown:
        raise ValueError(f"spec references unknown evidence ids: {unknown}")
    if has_summary and not spec.get("summary_evidence"):
        raise ValueError("summary_evidence must cite at least one supported fact")
    if spec.get("projects_note") and not spec.get("projects_note_evidence"):
        raise ValueError("projects_note must cite at least one evidence id")

    for item in _linkable_items(spec):
        for link in item.get("links", []):
            if not link.get("label") or not link.get("url", "").startswith("https://"):
                raise ValueError(f"link needs a label and an https url: {link}")

    for section in ("experience", "leadership", "projects", "skills"):
        for index, item in enumerate(spec.get(section, [])):
            if not item.get("evidence"):
                raise ValueError(f"{section}[{index}] must cite at least one evidence id")
            for bullet_index, bullet in enumerate(item.get("bullets", [])):
                if not bullet.get("evidence"):
                    raise ValueError(f"{section}[{index}].bullets[{bullet_index}] must cite evidence")
    entries = _education_entries(spec)
    if not entries:
        raise ValueError("education must contain at least one entry")
    for index, edu in enumerate(entries):
        if not edu.get("evidence"):
            raise ValueError(f"education[{index}] must cite at least one evidence id")

    for text in _iter_text(spec):
        if "—" in text or "–" in text:
            raise ValueError(f"non-ASCII dash found in {spec_path.name}: {text!r}")


def _set_font(run, st: Style, size: Pt | None = None, *, color=INK, bold=False, italic=False):
    run.font.name = st.font_name
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), st.font_name)
    rfonts.set(qn("w:hAnsi"), st.font_name)
    run.font.size = size
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic
    return run


def _run(p, st: Style, text: str, *, offset=0.0, color=INK, bold=False, italic=False):
    return _set_font(p.add_run(text), st, st.font(offset), color=color, bold=bold, italic=italic)


def _paragraph(doc: Document, st: Style, *, align=None, before=0.0, after=0.0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = st.gap(before)
    p.paragraph_format.space_after = st.gap(after)
    p.paragraph_format.line_spacing = st.line_spacing
    return p


def _bottom_border(p, *, color: str, size: str = "6") -> None:
    ppr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    ppr.append(pbdr)


def _section(doc: Document, st: Style, label: str) -> None:
    if st.technical:
        p = _paragraph(doc, st, before=8.0, after=2.6)
        p.paragraph_format.keep_with_next = True
        run = _run(p, st, label, offset=-0.3, bold=True)
        tracking = OxmlElement("w:spacing")
        tracking.set(qn("w:val"), "10")
        run._element.get_or_add_rPr().append(tracking)
        _bottom_border(p, color=RULE, size="4")
        return
    p = _paragraph(doc, st, before=6.2, after=2.0)
    _run(p, st, label, offset=-0.15, color=st.accent, bold=True)
    _bottom_border(p, color=st.accent_hex.upper())


def _right_tab(p, st: Style) -> None:
    p.paragraph_format.tab_stops.add_tab_stop(Inches(8.5 - 2 * st.margin_in), WD_TAB_ALIGNMENT.RIGHT)


def _item_header(doc: Document, st: Style, item: dict[str, Any]) -> None:
    p = _paragraph(doc, st, before=4.2, after=0.8)
    _right_tab(p, st)
    _run(p, st, item["title"], offset=0.65, bold=True)
    if item.get("organization"):
        _run(p, st, "  |  ", color=FAINT)
        _run(p, st, item["organization"], italic=True, color=MUTED)
    _run(p, st, "\t")
    right = " | ".join(part for part in (item.get("location"), item.get("dates")) if part)
    _run(p, st, right, offset=-0.9, color=FAINT)


def _hyperlink(p, st: Style, text: str, url: str, *, offset=0.0) -> None:
    """Add a clickable link that stays ordinary selectable text in the PDF."""
    rel_id = p.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    link.append(_run(p, st, text, offset=offset, color=LINK)._element)
    p._p.append(link)


def _rich(p, st: Style, text: str) -> None:
    """Render ``**emphasis**`` spans in running text as bold."""
    for index, part in enumerate(text.split("**")):
        if part:
            _run(p, st, part, bold=index % 2 == 1)


def _technical_entry(doc: Document, st: Style, item: dict[str, Any]) -> None:
    p = _paragraph(doc, st, before=4.0, after=0.5)
    p.paragraph_format.keep_with_next = True
    _right_tab(p, st)
    _run(p, st, item["title"], offset=0.5, bold=True)
    detail = item.get("organization") or item.get("subtitle")
    if detail:
        _run(p, st, "  |  ", color=FAINT)
        _run(p, st, detail, color=MUTED)
    _run(p, st, "\t")
    links = item.get("links", [])
    for index, link in enumerate(links):
        if index:
            _run(p, st, "  |  ", offset=-0.7, color=FAINT)
        _hyperlink(p, st, link["label"], link["url"], offset=-0.7)
    if item.get("dates"):
        if links:
            _run(p, st, "  |  ", offset=-0.7, color=FAINT)
        _run(p, st, item["dates"], offset=-0.5, color=MUTED)
    if item.get("description"):
        p = _paragraph(doc, st, after=0.5)
        p.paragraph_format.keep_with_next = True
        _run(p, st, item["description"])


def _bullet(doc: Document, st: Style, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.2 if st.technical else 0.24)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    gap = st.gap(0.7 if st.technical else 0.55)
    p.paragraph_format.space_before = gap
    p.paragraph_format.space_after = gap
    p.paragraph_format.line_spacing = st.line_spacing
    if st.technical:
        _rich(p, st, text)
    else:
        _run(p, st, text)


def _render_items(doc: Document, st: Style, items: list[dict[str, Any]]) -> None:
    for item in items:
        (_technical_entry if st.technical else _item_header)(doc, st, item)
        for bullet in item.get("bullets", []):
            _bullet(doc, st, bullet["text"])


def _render_summary(doc: Document, st: Style, section: dict[str, Any]) -> None:
    _section(doc, st, section["title"])
    p = _paragraph(doc, st, before=1.4, after=1.4)
    _run(p, st, section["spec"]["summary"])


def _render_alignment(doc: Document, st: Style, section: dict[str, Any]) -> None:
    _section(doc, st, section["title"])
    p = _paragraph(doc, st, before=1.4, after=1.4)
    _run(p, st, " | ".join(section["spec"]["alignment"]))


def _render_experience(doc: Document, st: Style, section: dict[str, Any]) -> None:
    _section(doc, st, section["title"])
    _render_items(doc, st, section["spec"]["experience"])


def _render_leadership(doc: Document, st: Style, section: dict[str, Any]) -> None:
    _section(doc, st, section["title"])
    _render_items(doc, st, section["spec"]["leadership"])


def _render_projects(doc: Document, st: Style, section: dict[str, Any]) -> None:
    _section(doc, st, section["title"])
    spec = section["spec"]
    if st.technical:
        if spec.get("projects_note"):
            p = _paragraph(doc, st, before=1.0, after=0.5)
            _run(p, st, spec["projects_note"], offset=-0.5, italic=True, color=MUTED)
        _render_items(doc, st, spec["projects"])
        return
    for item in spec["projects"]:
        p = _paragraph(doc, st, before=2.6, after=1.5)
        _run(p, st, f'{item["title"]}: ', bold=True)
        _run(p, st, item["description"])
        if item.get("tools"):
            _run(p, st, f'  {item["tools"]}', italic=True, color=MUTED)


def _skill_lines(doc: Document, st: Style, groups: list[dict[str, Any]]) -> None:
    first, rest = (1.2, 0.6) if st.technical else (1.5, 0.65)
    for i, group in enumerate(groups):
        p = _paragraph(doc, st, before=first if i == 0 else rest, after=rest)
        _run(p, st, f'{group["label"]}: ', bold=True)
        _run(p, st, group["items"])


def _render_skills(doc: Document, st: Style, section: dict[str, Any]) -> None:
    _section(doc, st, section["title"])
    _skill_lines(doc, st, section["spec"]["skills"])


def _technical_education(doc: Document, st: Style, edu: dict[str, Any]) -> None:
    p = _paragraph(doc, st, before=2.0, after=0.5)
    _right_tab(p, st)
    _run(p, st, edu["degree"], offset=0.5, bold=True)
    if edu.get("school"):
        _run(p, st, "  |  ", color=FAINT)
        _run(p, st, edu["school"], color=MUTED)
    if edu.get("date"):
        _run(p, st, "\t")
        _run(p, st, edu["date"], offset=-0.5, color=MUTED)
    if edu.get("detail"):
        p = _paragraph(doc, st, before=0.4)
        _run(p, st, edu["detail"], offset=-0.5, color=MUTED)


def _render_education(doc: Document, st: Style, section: dict[str, Any]) -> None:
    _section(doc, st, section["title"])
    if st.technical:
        for edu in _education_entries(section["spec"]):
            _technical_education(doc, st, edu)
        return
    for i, edu in enumerate(_education_entries(section["spec"])):
        p = _paragraph(doc, st, before=3.2 if i == 0 else 2.2, after=0.7)
        _right_tab(p, st)
        _run(p, st, edu["degree"], offset=0.55, bold=True)
        if edu.get("school"):
            _run(p, st, "  |  ", color=FAINT)
            _run(p, st, edu["school"], italic=True, color=MUTED)
        right = " | ".join(part for part in (edu.get("location"), edu.get("date")) if part)
        if right:
            _run(p, st, "\t")
            _run(p, st, right, offset=-0.9, color=FAINT)
        if edu.get("detail"):
            p = _paragraph(doc, st, before=0.5, after=0.5)
            _run(p, st, edu["detail"], offset=-0.15, color=MUTED)


BUILTIN_RENDERERS = {
    "summary": _render_summary,
    "alignment": _render_alignment,
    "experience": _render_experience,
    "leadership": _render_leadership,
    "projects": _render_projects,
    "skills": _render_skills,
    "education": _render_education,
}


# --- Custom sections ---------------------------------------------------------
# A custom section is {id, title, type, items}. Each type pairs a renderer
# (doc, style, section) with a validator (section) that raises ValueError with
# a path such as "sections[awards].items[0]".


def _validate_item(item: Any, path: str, fields: tuple[str, ...]) -> None:
    if not isinstance(item, dict):
        raise ValueError(f"{path} must be an object")
    for key in fields:
        if not isinstance(item.get(key), str) or not item[key].strip():
            raise ValueError(f"{path}.{key} must be a nonempty string")
    evidence = item.get("evidence")
    if not isinstance(evidence, list) or not evidence or not all(isinstance(e, str) and e for e in evidence):
        raise ValueError(f"{path} must cite at least one evidence id")


def _custom_items(section: dict[str, Any], fields: tuple[str, ...]) -> list[dict[str, Any]]:
    path = f'sections[{section["id"]}]'
    items = section.get("items")
    if not isinstance(items, list) or not items:
        raise ValueError(f"{path}.items must be a nonempty list")
    for index, item in enumerate(items):
        _validate_item(item, f"{path}.items[{index}]", fields)
    return items


def _validate_bullets(section: dict[str, Any]) -> None:
    _custom_items(section, ("text",))


def _validate_skills_list(section: dict[str, Any]) -> None:
    _custom_items(section, ("label", "items"))


def _validate_entries(section: dict[str, Any]) -> None:
    for index, item in enumerate(_custom_items(section, ("title",))):
        path = f'sections[{section["id"]}].items[{index}]'
        for key in ("organization", "location", "dates"):
            if key in item and not isinstance(item[key], str):
                raise ValueError(f"{path}.{key} must be a string")
        bullets = item.get("bullets", [])
        if not isinstance(bullets, list):
            raise ValueError(f"{path}.bullets must be a list")
        for bullet_index, bullet in enumerate(bullets):
            _validate_item(bullet, f"{path}.bullets[{bullet_index}]", ("text",))


def _render_bullets(doc: Document, st: Style, section: dict[str, Any]) -> None:
    _section(doc, st, section["title"])
    for item in section["items"]:
        _bullet(doc, st, item["text"])


def _render_skills_list(doc: Document, st: Style, section: dict[str, Any]) -> None:
    _section(doc, st, section["title"])
    _skill_lines(doc, st, section["items"])


def _render_entries(doc: Document, st: Style, section: dict[str, Any]) -> None:
    _section(doc, st, section["title"])
    _render_items(doc, st, section["items"])


SECTION_TYPES = {
    "bullets": (_render_bullets, _validate_bullets),
    "skills-list": (_render_skills_list, _validate_skills_list),
    "entries": (_render_entries, _validate_entries),
}


def _custom_sections(spec: dict[str, Any]) -> dict[str, dict[str, Any]]:
    custom = spec.get("sections", [])
    if not isinstance(custom, list):
        raise ValueError("sections must be a list")
    sections: dict[str, dict[str, Any]] = {}
    for index, section in enumerate(custom):
        if not isinstance(section, dict):
            raise ValueError(f"sections[{index}] must be an object")
        sid = section.get("id")
        if not isinstance(sid, str) or not CUSTOM_ID.fullmatch(sid) or sid in SECTION_TITLES:
            raise ValueError(
                f"invalid custom section id: {sid!r} (use [a-z0-9_-]+, not a built-in id)"
            )
        if sid in sections:
            raise ValueError(f"duplicate section id: {sid}")
        title = section.get("title")
        if not isinstance(title, str) or not title.strip():
            raise ValueError(f"sections[{sid}].title must be a nonempty string")
        kind = section.get("type")
        if not isinstance(kind, str) or kind not in SECTION_TYPES:
            raise ValueError(f"unknown section type for sections[{sid}]: {kind!r} (known: {sorted(SECTION_TYPES)})")
        SECTION_TYPES[kind][1](section)
        sections[sid] = section
    return sections


def _section_order(spec: dict[str, Any], custom: dict[str, dict[str, Any]]) -> list[str]:
    layout = layout_of(spec)
    if "section_order" not in spec:
        default = list(layout.order)
        cut = default.index(layout.custom_before) if layout.custom_before else len(default)
        return default[:cut] + list(custom) + default[cut:]
    order = spec["section_order"]
    if not isinstance(order, list) or not all(isinstance(sid, str) for sid in order):
        raise ValueError("section_order must be a list of section ids")
    duplicates = sorted({sid for sid in order if order.count(sid) > 1})
    if duplicates:
        raise ValueError(f"duplicate ids in section_order: {duplicates}")
    unknown = sorted(set(order) - set(layout.titles) - set(custom))
    if unknown:
        raise ValueError(f"unknown section ids in section_order: {unknown}")
    required = {"experience", "education"} | set(custom)
    if spec.get("skills"):
        required.add("skills")
    missing = sorted(required - set(order))
    if missing:
        raise ValueError(f"section_order missing required sections: {missing}")
    return order


def _resolve_sections(spec: dict[str, Any]) -> list[dict[str, Any]]:
    """Validate layout fields and return the sections to render, in order.

    Built-in entries are {id, type, title, spec}; custom entries are the
    custom section dict with its title resolved.
    """
    layout = layout_of(spec)
    custom = _custom_sections(spec)
    order = _section_order(spec, custom)
    titles = spec.get("section_titles", {})
    if not isinstance(titles, dict):
        raise ValueError("section_titles must be an object of section id -> heading")
    unknown = sorted(set(titles) - set(layout.titles) - set(custom))
    if unknown:
        raise ValueError(f"unknown section_titles keys: {unknown}")
    blank = sorted(k for k, v in titles.items() if not isinstance(v, str) or not v.strip())
    if blank:
        raise ValueError(f"section_titles values must be nonempty strings: {blank}")

    resolved = []
    for sid in order:
        if sid in OPTIONAL_SECTIONS and not spec.get(sid):
            continue
        if sid in custom:
            section = dict(custom[sid])
        else:
            section = {"id": sid, "type": sid, "title": layout.titles[sid], "spec": spec}
        section["title"] = titles.get(sid, section["title"]).strip().upper()
        resolved.append(section)
    return resolved


def section_headings(spec: dict[str, Any]) -> list[str]:
    """Rendered section headings (uppercase) in layout order."""
    return [section["title"] for section in _resolve_sections(spec)]


def _classic_header(doc: Document, st: Style, spec: dict[str, Any], person: dict[str, str]) -> None:
    p = _paragraph(doc, st, align=WD_ALIGN_PARAGRAPH.CENTER, after=0.8)
    _run(p, st, person["name"], offset=13.0, bold=True)

    if spec.get("headline"):
        p = _paragraph(doc, st, align=WD_ALIGN_PARAGRAPH.CENTER, after=1.5)
        _run(p, st, spec["headline"], offset=0.1, color=st.accent, bold=True)

    p = _paragraph(doc, st, align=WD_ALIGN_PARAGRAPH.CENTER, after=2.8)
    _run(p, st, " | ".join(contact_fields(spec, person).values()), offset=-1.1, color=MUTED)


def _technical_header(doc: Document, st: Style, spec: dict[str, Any], person: dict[str, str]) -> None:
    p = _paragraph(doc, st, after=2.0)
    _run(p, st, person["name"], offset=11.0, bold=True)

    p = _paragraph(doc, st, after=1.0)
    links = contact_links(person)
    for index, (key, value) in enumerate(contact_fields(spec, person).items()):
        if index:
            _run(p, st, "  |  ", offset=-0.5, color=FAINT)
        if key in links:
            _hyperlink(p, st, value, links[key], offset=-0.5)
        else:
            _run(p, st, value, offset=-0.5, color=MUTED)


def build_document(inputs: ResumeInputs, style: Style) -> Document:
    spec = inputs.spec
    person = inputs.person
    layout = layout_of(spec)
    style = replace(style, line_spacing=layout.line_spacing, technical=layout.technical)

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, side, Inches(style.margin_in))

    normal = doc.styles["Normal"]
    normal.font.name = style.font_name
    normal.font.size = style.font()
    normal.font.color.rgb = INK
    normal.paragraph_format.line_spacing = style.line_spacing

    if layout.technical:
        doc.core_properties.author = person["name"]
        doc.core_properties.title = f'{person["name"]} - Resume'
        _technical_header(doc, style, spec, person)
    else:
        _classic_header(doc, style, spec, person)

    for block in _resolve_sections(spec):
        renderer = BUILTIN_RENDERERS.get(block["id"]) or SECTION_TYPES[block["type"]][0]
        renderer(doc, style, block)

    return doc


def build_and_save(inputs: ResumeInputs, style: Style, output_path: Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    build_document(inputs, style).save(output_path)
    return output_path
