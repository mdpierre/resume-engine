"""Technical layout tests: projects first, no summary, real hyperlinks.

    python -m unittest -v tests/test_technical.py
"""

from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import resume_checks as rc  # noqa: E402
from resume_factory import (  # noqa: E402
    LAYOUTS,
    build_and_save,
    expected_links,
    load_inputs,
    section_headings,
)

EXAMPLE = ROOT / "example"
FACTS = EXAMPLE / "resume_facts.json"
BASE_SPEC = EXAMPLE / "specs" / "base.json"
TECHNICAL_SPEC = EXAMPLE / "specs" / "technical.json"
HEADINGS = [
    "EDUCATION",
    "TECHNICAL SKILLS",
    "SELECTED TECHNICAL PROJECTS",
    "PROFESSIONAL EXPERIENCE",
    "LEADERSHIP & ACTIVITIES",
]


def _has_tools() -> bool:
    try:
        rc.find_soffice()
        rc.find_pdftotext()
        rc.find_pdfinfo()
    except FileNotFoundError:
        return False
    return True


class TechnicalLayoutTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name)
        self.spec = json.loads(TECHNICAL_SPEC.read_text(encoding="utf-8"))
        self.spec.pop("sections")

    def inputs(self, **changes):
        path = self.root / "spec.json"
        path.write_text(json.dumps({**self.spec, **changes}), encoding="utf-8")
        return load_inputs(path, FACTS)

    def render(self, **changes):
        inputs = self.inputs(**changes)
        return inputs, build_and_save(inputs, inputs.base_style(), self.root / "resume.docx")

    def text(self, docx):
        return "\n".join(p.text for p in Document(str(docx)).paragraphs)

    def test_classic_is_the_default(self):
        spec = json.loads(BASE_SPEC.read_text(encoding="utf-8"))
        self.assertNotIn("layout", spec)
        self.assertEqual(section_headings(spec)[0], "PROFILE")
        self.assertEqual(LAYOUTS["classic"].titles["summary"], "PROFILE")

    def test_order_has_no_profile_or_alignment(self):
        inputs, docx = self.render(headline="Should Not Render")
        self.assertEqual(section_headings(inputs.spec), HEADINGS)
        self.assertEqual(rc.docx_headings(docx), HEADINGS)
        text = self.text(docx)
        for removed in ("PROFILE", "ROLE ALIGNMENT", "Should Not Render"):
            self.assertNotIn(removed, text)

    def test_variant_can_reorder_and_rename(self):
        order = ["education", "skills", "experience", "projects", "leadership"]
        inputs, docx = self.render(section_order=order, section_titles={"projects": "Projects"})
        self.assertEqual(
            rc.docx_headings(docx),
            ["EDUCATION", "TECHNICAL SKILLS", "PROFESSIONAL EXPERIENCE", "PROJECTS", "LEADERSHIP & ACTIVITIES"],
        )
        self.assertEqual(section_headings(inputs.spec), rc.docx_headings(docx))

    def test_custom_sections_render_after_the_layout_order(self):
        section = {
            "id": "results",
            "title": "Selected Results",
            "type": "bullets",
            "items": [{"text": "Cut report prep to **under 1 hour**.", "evidence": ["experience.lakeshore_dashboard"]}],
        }
        _, docx = self.render(sections=[section])
        self.assertEqual(rc.docx_headings(docx), HEADINGS + ["SELECTED RESULTS"])
        self.assertIn("Cut report prep to under 1 hour.", self.text(docx))

    def test_rejects_unknown_layout_section_and_insecure_link(self):
        bad_link = {**self.spec["projects"][0], "links": [{"label": "x", "url": "http://x.test"}]}
        no_label = {**self.spec["projects"][0], "links": [{"url": "https://x.test"}]}
        cases = (
            ({"layout": "fancy"}, "unknown layout"),
            ({"section_order": ["education", "experience", "skills", "summary"]}, "unknown section"),
            ({"section_titles": {"summary": "Summary"}}, "unknown section_titles"),
            ({"projects": [bad_link]}, "https"),
            ({"projects": [no_label]}, "label"),
        )
        for changes, message in cases:
            with self.subTest(message=message), self.assertRaisesRegex(ValueError, message):
                self.inputs(**changes)

    def test_project_note_must_cite_known_evidence(self):
        with self.assertRaisesRegex(ValueError, "projects_note"):
            self.inputs(projects_note="Note.", projects_note_evidence=[])
        with self.assertRaisesRegex(ValueError, "unknown evidence"):
            self.inputs(projects_note="Note.", projects_note_evidence=["unsupported.fact"])

    def test_contact_and_project_links_are_real_hyperlinks(self):
        inputs, docx = self.render()
        rels = Document(str(docx)).part.rels.values()
        targets = {r.target_ref for r in rels if r.is_external}
        self.assertEqual(targets, set(expected_links(inputs)))
        self.assertIn("mailto:jordan.rivera@example.com", targets)
        self.assertIn("https://www.linkedin.com/in/jrivera/", targets)
        self.assertIn("https://github.com/jrivera", targets)
        self.assertIn("https://example.com/jrivera/transit-delay-tracker", targets)

    def test_links_only_expected_for_technical_and_visible_contacts(self):
        self.assertEqual(expected_links(load_inputs(BASE_SPEC, FACTS)), [])
        inputs, docx = self.render(hide_contact=["github"], hide_phone=True)
        self.assertNotIn("https://github.com/jrivera", expected_links(inputs))
        text = self.text(docx)
        self.assertNotIn("github.com/jrivera", text)
        self.assertNotIn(inputs.person["phone"], text)

    def test_document_metadata_names_the_candidate(self):
        inputs, docx = self.render()
        self.assertEqual(Document(str(docx)).core_properties.author, inputs.person["name"])

    def test_inline_bold_markup_never_reaches_the_page(self):
        _, docx = self.render()
        paragraphs = Document(str(docx)).paragraphs
        self.assertNotIn("**", "\n".join(p.text for p in paragraphs))
        self.assertTrue(any(run.bold and run.text == "8 volunteers" for p in paragraphs for run in p.runs))

    def test_section_order_rejects_an_unexpected_known_heading(self):
        text = "Jordan Rivera\njordan.rivera@example.com\nEDUCATION\nEXPERIENCE\nPROJECTS"
        with patch.object(rc, "pdf_text", return_value=text):
            self.assertTrue(rc.check_section_order(Path("unused.pdf"), ["EDUCATION", "EXPERIENCE", "PROJECTS"]).passed)
            self.assertFalse(rc.check_section_order(Path("unused.pdf"), ["EDUCATION", "EXPERIENCE"]).passed)

    def test_text_layer_check_counts_words(self):
        words = [rc.Word(0, 0, 1, 1)] * rc.MIN_TEXT_WORDS
        self.assertTrue(rc.check_text_layer([rc.Page(612, 792, words)]).passed)
        self.assertFalse(rc.check_text_layer([rc.Page(612, 792, words[1:])]).passed)

    @unittest.skipUnless(_has_tools(), "LibreOffice/Poppler not installed")
    def test_rendered_pdf_passes_heading_link_and_text_checks(self):
        inputs, docx = self.render()
        pdf = rc.render_pdf(docx, self.root)
        links = expected_links(inputs)
        report = rc.evaluate(pdf, expected_headings=section_headings(inputs.spec), expected_links=links)
        by_name = {check.name: check for check in report.checks}
        for name in ("section_order", "selectable_text", "links"):
            self.assertTrue(by_name[name].passed, by_name[name].detail)
        self.assertFalse(rc.check_links(pdf, links + ["https://example.com/missing"]).passed)


if __name__ == "__main__":
    unittest.main(verbosity=2)
